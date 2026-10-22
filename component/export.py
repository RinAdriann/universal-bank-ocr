import io
import re
from dataclasses import asdict
from datetime import date, datetime
from typing import Any

import pandas as pd
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from .constants import DETAIL_COLS, MONTH_NAMES_ID, REVIEW_COLS
from .transactioncat import category


def safe_str(val: Any) -> str:
    text = "" if val is None else str(val)
    if text.startswith(("=", "+", "-", "@")):
        return f"'{text}"
    return text


def format_date(dt: Any) -> str:
    if dt is None:
        return ""
    if isinstance(dt, (date, datetime)):
        return dt.strftime("%d/%m/%Y")
    return str(dt)


def transaction_record(row: Any, format_dt: bool = False, escape_excel: bool = True) -> dict[str, Any]:
    desc = safe_str(row.description) if escape_excel else (row.description or "")
    dt_val = format_date(row.date) if format_dt else (row.date if row.date is not None else "")
    return {
        "Tanggal": dt_val,
        "Kredit": row.credit,
        "Debit": row.debit,
        "Saldo": row.balance,
        "Deskripsi": desc,
        "Kategori Transaksi": category(row.description, row.debit, row.credit),
    }


def rows_df(rows: list[Any], format_dt: bool = True, escape_excel: bool = False) -> pd.DataFrame:
    records = [transaction_record(r, format_dt=format_dt, escape_excel=escape_excel) for r in rows]
    return pd.DataFrame(records, columns=DETAIL_COLS)


def review_df(rows: list[Any], detailed: bool = True, escape_excel: bool = False) -> pd.DataFrame:
    records = []
    for row in rows:
        raw = safe_str(row.raw_text) if escape_excel else (row.raw_text or "")
        flags = ", ".join(row.validation)
        if detailed:
            records.append({
                "Halaman": row.page,
                "Confidence": round(getattr(row, "confidence", 1.0), 3),
                "Human review status": "Needs review" if row.validation else "Pending review",
                "Validation flags": flags,
                "Full detected text": raw,
            })
        else:
            records.append({
                "Halaman": row.page,
                "Validation flags": flags,
                "Full detected text": raw,
            })
    cols = None if detailed else REVIEW_COLS
    return pd.DataFrame(records, columns=cols)


def serialize_row(row: Any) -> dict[str, Any]:
    item = asdict(row)
    item["date"] = row.date.isoformat() if isinstance(row.date, (date, datetime)) else row.date
    item.pop("date_parts", None)
    item.pop("time", None)
    item.pop("pending_dir", None)
    return item


def month_key(dt: Any) -> tuple[int, int] | None:
    if isinstance(dt, (date, datetime)):
        return (dt.year, dt.month)
    if isinstance(dt, str):
        match = re.match(r"^(\d{2})/(\d{2})(?:/(\d{4}))?$", dt)
        if match:
            return (int(match.group(3) or 0), int(match.group(2)))
    return None


def month_label(ym: tuple[int, int]) -> str:
    year, month = ym
    name = MONTH_NAMES_ID.get(month, f"Bulan {month}")
    return f"{name} {year}" if year else name


def saldo_mengendap(rows: list[Any]) -> dict[tuple[int, int], list[tuple[Any, float]]]:
    by_day: dict[Any, list[Any]] = {}
    for row in rows:
        if getattr(row, "is_opening", False) or row.date is None or row.balance is None:
            continue
        if month_key(row.date):
            by_day.setdefault(row.date, []).append(row)

    by_month: dict[tuple[int, int], list[tuple[Any, float]]] = {}
    for dt, day_rows in by_day.items():
        ym = month_key(dt)
        if not ym:
            continue
        if any(getattr(r, "time", None) is not None for r in day_rows):
            last_row = max(
                enumerate(day_rows),
                key=lambda item: (getattr(item[1], "time", None) or (-1, -1, -1), item[0]),
            )[1]
        else:
            last_row = day_rows[-1]
        by_month.setdefault(ym, []).append((dt, float(last_row.balance)))

    for ym in by_month:
        by_month[ym].sort(key=lambda pair: pair[0] if isinstance(pair[0], date) else date.min)
    return dict(sorted(by_month.items()))


def top5_monthly(rows: list[Any]) -> dict[tuple[int, int], dict[str, list[Any]]]:
    by_month: dict[tuple[int, int], list[Any]] = {}
    for row in rows:
        if getattr(row, "is_opening", False) or row.date is None:
            continue
        ym = month_key(row.date)
        if ym:
            by_month.setdefault(ym, []).append(row)

    out: dict[tuple[int, int], dict[str, list[Any]]] = {}
    for ym in sorted(by_month):
        month_rows = by_month[ym]
        credits = [r for r in month_rows if r.credit and r.credit > 0]
        debits = [r for r in month_rows if r.debit and r.debit > 0]
        top_cr = sorted(credits, key=lambda r: float(r.credit), reverse=True)[:5]
        top_db = sorted(debits, key=lambda r: float(r.debit), reverse=True)[:5]
        out[ym] = {"cr": top_cr, "db": top_db}
    return out


def sheet_title(prefix: str, bank: str, acc: str, used: set[str]) -> str:
    clean_bank = re.sub(r"[\[\]:*?/\\]", "", (bank or "").strip())
    clean_acc = re.sub(r"[\[\]:*?/\\]", "", (acc or "").strip())
    if clean_acc in ("Default", "Unknown"):
        clean_acc = ""

    if clean_bank and clean_acc:
        spaced = f"{clean_bank} - {clean_acc}"
        compact = f"{clean_bank}-{clean_acc}"
    else:
        spaced = compact = clean_bank or clean_acc

    if not spaced:
        name = prefix
    else:
        name = f"{prefix} {spaced}"
        if len(name) > 31:
            name = f"{prefix} {compact}"
        if len(name) > 31:
            short_prefix = "Saldo" if prefix.lower().startswith("saldo") else prefix[: max(2, 30 - len(compact))].rstrip()
            name = f"{short_prefix} {compact}"

    name = name[:31]
    base = name
    idx = 2
    while name in used:
        tag = f"_{idx}"
        name = f"{base[: 31 - len(tag)]}{tag}"
        idx += 1
    used.add(name)
    return name


def write_date_cell(ws: Any, row_idx: int, col_idx: int, dt: Any) -> None:
    cell = ws.cell(row_idx, col_idx, dt if dt is not None else "")
    if isinstance(dt, (date, datetime)):
        cell.number_format = "dd/mm/yyyy"


def write_money_cell(ws: Any, row_idx: int, col_idx: int, val: float | None) -> None:
    cell = ws.cell(row_idx, col_idx, val)
    cell.number_format = "#,##0.00"


def saldo_mengendap_sheet(ws: Any, rows: list[Any], bold_font: Font) -> None:
    monthly = saldo_mengendap(rows)
    if not monthly:
        ws.cell(1, 1, "Tanggal").font = bold_font
        ws.cell(1, 2, "Saldo Mengendap").font = bold_font
        ws.column_dimensions["A"].width = 14
        ws.column_dimensions["B"].width = 18
        return

    for idx, (ym, items) in enumerate(monthly.items()):
        col_date = 1 + idx * 3
        col_bal = col_date + 1
        ws.cell(1, col_date, month_label(ym)).font = bold_font
        ws.merge_cells(start_row=1, start_column=col_date, end_row=1, end_column=col_bal)
        ws.cell(2, col_date, "Tanggal").font = bold_font
        ws.cell(2, col_bal, "Saldo Mengendap").font = bold_font

        for row_idx, (dt, bal) in enumerate(items, start=3):
            write_date_cell(ws, row_idx, col_date, dt)
            write_money_cell(ws, row_idx, col_bal, bal)

        ws.column_dimensions[get_column_letter(col_date)].width = 14
        ws.column_dimensions[get_column_letter(col_bal)].width = 18
        ws.column_dimensions[get_column_letter(col_bal + 1)].width = 4
    ws.freeze_panes = "A3"


def write_top5_block(
    ws: Any,
    start_row: int,
    col_start: int,
    title: str,
    amt_header: str,
    rows: list[Any],
    field: str,
    bold_font: Font,
) -> None:
    ws.cell(start_row, col_start, title).font = bold_font
    ws.merge_cells(start_row=start_row, start_column=col_start, end_row=start_row, end_column=col_start + 2)
    for offset, header in enumerate(("Tanggal", amt_header, "Deskripsi")):
        ws.cell(start_row + 1, col_start + offset, header).font = bold_font

    for idx, row in enumerate(rows, start=start_row + 2):
        write_date_cell(ws, idx, col_start, row.date)
        write_money_cell(ws, idx, col_start + 1, getattr(row, field))
        ws.cell(idx, col_start + 2, safe_str(row.description))


def top5_sheet(ws: Any, rows: list[Any], bold_font: Font) -> None:
    monthly = top5_monthly(rows)
    if not monthly:
        ws.cell(1, 1, "Tanggal").font = bold_font
        ws.cell(1, 2, "Nominal").font = bold_font
        ws.cell(1, 3, "Deskripsi").font = bold_font
        return

    for idx, (ym, data) in enumerate(monthly.items()):
        col_start = 1 + idx * 4
        label = month_label(ym)
        write_top5_block(ws, 1, col_start, f"{label} - Top 5 CR", "Kredit", data["cr"], "credit", bold_font)
        write_top5_block(ws, 9, col_start, f"{label} - Top 5 DB", "Debit", data["db"], "debit", bold_font)

        ws.column_dimensions[get_column_letter(col_start)].width = 14
        ws.column_dimensions[get_column_letter(col_start + 1)].width = 18
        ws.column_dimensions[get_column_letter(col_start + 2)].width = 50
        ws.column_dimensions[get_column_letter(col_start + 3)].width = 4
    ws.freeze_panes = "A3"


def excel_bytes(result: dict[str, Any]) -> bytes:
    output = io.BytesIO()
    rows = result.get("rows", [])
    accounts_map = result.get("accounts")

    if accounts_map:
        groups = [
            (data.get("account") or "", data.get("bank") or "", data["rows"])
            for data in accounts_map.values()
        ]
    else:
        groups = [("", "", rows)]

    used_names: set[str] = set()
    bold_font = Font(bold=True)
    sheet_specs = (
        ("Detail", lambda _: True),
        ("In", lambda r: bool(r.credit and r.credit > 0)),
        ("Out", lambda r: bool(r.debit and r.debit > 0)),
    )

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        for acc, bank, acc_rows in groups:
            for prefix, pred in sheet_specs:
                sname = sheet_title(prefix, bank, acc, used_names)
                filtered = [r for r in acc_rows if pred(r)]
                rows_df(filtered, format_dt=False, escape_excel=True).to_excel(writer, sheet_name=sname, index=False)

                ws = writer.sheets[sname]
                ws.freeze_panes = "A2"
                for col, width in zip(["A", "B", "C", "D", "E", "F"], [14, 16, 16, 16, 80, 22]):
                    ws.column_dimensions[col].width = width
                for excel_row in ws.iter_rows(min_row=2):
                    if isinstance(excel_row[0].value, (date, datetime)):
                        excel_row[0].number_format = "dd/mm/yyyy"
                    for cell in excel_row[1:4]:
                        cell.number_format = "#,##0.00"

            sm_name = sheet_title("Saldo Mengendap", bank, acc, used_names)
            saldo_mengendap_sheet(writer.book.create_sheet(title=sm_name), acc_rows, bold_font)

            t5_name = sheet_title("Top 5", bank, acc, used_names)
            top5_sheet(writer.book.create_sheet(title=t5_name), acc_rows, bold_font)

            rev_name = sheet_title("Review", bank, acc, used_names)
            review_df(acc_rows, detailed=False, escape_excel=True).to_excel(writer, sheet_name=rev_name, index=False)
            rev_ws = writer.sheets[rev_name]
            rev_ws.freeze_panes = "A2"
            for col, width in zip(["A", "B", "C"], [10, 28, 100]):
                rev_ws.column_dimensions[col].width = width

    return output.getvalue()


__all__ = [
    "excel_bytes",
    "format_date",
    "month_label",
    "review_df",
    "rows_df",
    "saldo_mengendap",
    "serialize_row",
    "top5_monthly",
]
